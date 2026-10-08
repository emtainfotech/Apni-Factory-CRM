import os, sys
import django

# Setup Django environment
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
os.environ.setdefault('DJANGO_SETTINGS_MODULE', 'ApniFactoryCRM.settings')
django.setup()

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from hostinger_data.models import MainCategories, Categories

os.makedirs('static/samples', exist_ok=True)
target_path = 'static/samples/shade_cards_sample_template.xlsx'

wb = openpyxl.Workbook()

# Sheet 1: Shade Cards
ws = wb.active
ws.title = 'Shade Cards'

headers = ['Name', 'Maincategory', 'Category', 'Hexcode', 'Image', 'Status', 'User', 'Adminmsg']
ws.append(headers)

sample_rows = [
    ['Hazel nut -8569', 'Paint & Distemper', 'Exterior Emulsion Water Based', '#fffdf1', '', 1, 'admin', 'Done'],
    ['scarlet -8085', 'Paint & Distemper', 'Exterior Emulsion Water Based', '#be9d7c', '', 1, 'admin', 'Done'],
    ['Silver Ash', 'Paint & Distemper', 'Hammertone Finish', '', 'shadecard/KK2h4P5m34LH4kWlzcFAAEux4omyW6-metaU2lsdmVyIEFzaCBIYW1tZXJ0b25lLnBuZw==-.png', 1, 'admin', 'Done'],
    ['Apricot illusion -7979', 'Paint & Distemper', 'Exterior Emulsion Water Based', '#fcefc5', '', 1, 'admin', 'Done'],
    ['Tree Of Life -7691', 'Paint & Distemper', 'Exterior Emulsion Water Based', '#d1e7c3', '', 1, 'admin', 'Done'],
]
for row in sample_rows:
    ws.append(row)

header_fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
data_font = Font(name='Segoe UI', size=10)
border_side = Side(border_style='thin', color='CBD5E1')
cell_border = Border(left=border_side, right=border_side, top=border_side, bottom=border_side)

for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

for row in ws.iter_rows(min_row=2, max_row=len(sample_rows) + 1, min_col=1, max_col=len(headers)):
    for cell in row:
        cell.font = data_font
        cell.border = cell_border
        if cell.column in [6, 7]:
            cell.alignment = Alignment(horizontal='center', vertical='center')

col_widths = {'A': 28, 'B': 25, 'C': 32, 'D': 16, 'E': 35, 'F': 12, 'G': 15, 'H': 20}
for col_letter, width in col_widths.items():
    ws.column_dimensions[col_letter].width = width

# Sheet 2: Categories Reference
ws_ref = wb.create_sheet(title='Categories Reference')
ref_headers = ['Main Category ID', 'Main Category Name', 'Category ID', 'Category Name (Subcategory)']
ws_ref.append(ref_headers)

ref_fill = PatternFill(start_color='0D9488', end_color='0D9488', fill_type='solid')
ref_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
for cell in ws_ref[1]:
    cell.fill = ref_fill
    cell.font = ref_font
    cell.alignment = Alignment(horizontal='center', vertical='center')

main_cats = {m.id: m.name for m in MainCategories.objects.all()}
categories = Categories.objects.all().order_by('maincategory_id', 'name')
for cat in categories:
    ws_ref.append([
        cat.maincategory_id,
        main_cats.get(cat.maincategory_id, 'Unknown'),
        cat.id,
        cat.name
    ])

for row in ws_ref.iter_rows(min_row=2, max_col=4):
    for cell in row:
        cell.font = data_font
        cell.border = cell_border

ws_ref.column_dimensions['A'].width = 20
ws_ref.column_dimensions['B'].width = 30
ws_ref.column_dimensions['C'].width = 18
ws_ref.column_dimensions['D'].width = 35

# Sheet 3: Instructions
ws_rules = wb.create_sheet(title='Instructions')
rules = [
    ['RULE & FIELD SPECIFICATIONS FOR BULK SHADE CARD IMPORT'],
    [''],
    ['1. Name (Mandatory)', 'Shade name (e.g. "Hazel nut -8569"). Max 100 characters.'],
    ['2. Maincategory (Mandatory)', 'Main category Name (e.g. "Paint & Distemper") or Main Category ID (e.g. 2). See "Categories Reference" sheet.'],
    ['3. Category (Mandatory)', 'Category Name (e.g. "Exterior Emulsion Water Based") or Category ID (e.g. 40). Must belong to selected Main Category.'],
    ['4. MANDATE: Hexcode OR Image', 'AT LEAST ONE must be provided. Both cannot be empty!'],
    ['   - Hexcode', 'Color hex code like #fffdf1 or #be9d7c. If "#" is omitted, it will be added automatically.'],
    ['   - Image', 'Filename, path or URL (e.g. "shadecard/image.png" or full URL).'],
    ['5. Status (Optional)', '1 for Active (Default), 0 for Inactive.'],
    ['6. User (Optional)', 'Username or User ID (e.g. "admin" or "1"). Default is admin (1).'],
    ['7. Adminmsg (Optional)', 'Short message or note (e.g. "Done"). Max 255 characters.']
]
for r in rules:
    ws_rules.append(r)

ws_rules.column_dimensions['A'].width = 35
ws_rules.column_dimensions['B'].width = 65

wb.save(target_path)
print('Successfully generated static sample Excel template at:', target_path, 'Size:', os.path.getsize(target_path))
